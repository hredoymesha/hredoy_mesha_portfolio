"""Offline import and export adapters for supported reference formats."""

from __future__ import annotations

import base64
import csv
import io
import json
import re
from dataclasses import dataclass
from html import escape
from typing import Any

from bibtexparser import loads as load_bibtex
from bibtexparser.bparser import BibTexParser
from docx import Document

from citecraft.engine import render
from citecraft.models import Contributor, ExportRequest, Reference, RenderRequest


class FormatError(ValueError):
    """An import or export payload cannot be interpreted safely."""


@dataclass(frozen=True)
class ExportedFile:
    filename: str
    content_type: str
    content: str
    encoding: str = "utf-8"


def import_references(format_name: str, content: str) -> list[Reference]:
    try:
        if format_name in {"citecraft-json", "csl-json"}:
            return _import_json(format_name, content)
        if format_name == "ris":
            rows = _parse_ris(content)
        elif format_name == "bibtex":
            rows = _parse_bibtex(content)
        elif format_name == "csv":
            rows = _parse_csv(content)
        else:
            raise FormatError(f"Unsupported import format: {format_name}")
        return [Reference.model_validate(row) for row in rows]
    except (ValueError, KeyError, TypeError) as error:
        if isinstance(error, FormatError):
            raise
        raise FormatError(f"Could not import {format_name}: {error}") from error


def export_references(request: ExportRequest) -> ExportedFile:
    format_name = request.format
    references = request.references
    if format_name == "citecraft-json":
        content = json.dumps(
            {
                "schema_version": 1,
                "references": [item.model_dump(mode="json") for item in references],
            },
            ensure_ascii=False,
            indent=2,
        )
        return ExportedFile("citecraft-project.json", "application/json", content)
    if format_name == "csl-json":
        content = json.dumps(
            [reference_to_csl_json(item) for item in references],
            ensure_ascii=False,
            indent=2,
        )
        return ExportedFile("references.csl.json", "application/json", content)
    if format_name == "bibtex":
        return ExportedFile("references.bib", "application/x-bibtex", _export_bibtex(references))
    if format_name == "ris":
        return ExportedFile(
            "references.ris", "application/x-research-info-systems", _export_ris(references)
        )
    if format_name == "csv":
        return ExportedFile("references.csv", "text/csv", _export_csv(references))

    rendered = render(RenderRequest(references=references, style=request.style))
    entries = rendered["bibliography"]
    if format_name == "plain-text":
        return ExportedFile("bibliography.txt", "text/plain", "\n".join(entries))
    if format_name == "markdown":
        return ExportedFile(
            "bibliography.md",
            "text/markdown",
            "\n\n".join(f"{index}. {entry}" for index, entry in enumerate(entries, start=1)),
        )
    if format_name == "html":
        html = "<ol>\n" + "\n".join(f"  <li>{escape(entry)}</li>" for entry in entries) + "\n</ol>"
        return ExportedFile("bibliography.html", "text/html", html)
    if format_name == "docx":
        return _export_docx(entries)
    raise FormatError(f"Unsupported export format: {format_name}")


def reference_to_csl_json(reference: Reference) -> dict[str, Any]:
    from citecraft.engine import reference_to_csl

    return reference_to_csl(reference)


def _import_json(format_name: str, content: str) -> list[Reference]:
    payload = json.loads(content)
    if format_name == "citecraft-json":
        if not isinstance(payload, dict) or payload.get("schema_version") != 1:
            raise FormatError("Unsupported CiteCraft JSON version.")
        rows = payload.get("references")
    else:
        rows = payload if isinstance(payload, list) else [payload]
        rows = [_from_csl_json(item) for item in rows]
    if not isinstance(rows, list) or not rows:
        raise FormatError("The import file contains no references.")
    return [Reference.model_validate(row) for row in rows]


def _from_csl_json(item: Any) -> dict[str, Any]:
    if not isinstance(item, dict):
        raise FormatError("Each CSL JSON item must be an object.")
    source_type = item.get("type")
    supported_types = {
        "article-journal": "article-journal",
        "book": "book",
        "chapter": "chapter",
        "webpage": "webpage",
        "report": "report",
        "thesis": "thesis",
        "article": "article-journal",
        "paper-conference": "article-journal",
        "entry-encyclopedia": "chapter",
        "entry-dictionary": "chapter",
    }
    if source_type not in supported_types:
        raise FormatError(f"Unsupported CSL JSON type: {source_type!r}")
    issued = item.get("issued") or {}
    date_parts = issued.get("date-parts", [])
    pubdate = (
        "-".join(
            str(part).zfill(2) if index else str(part) for index, part in enumerate(date_parts[0])
        )
        if date_parts
        else None
    )
    authors = [_from_csl_contributor(person) for person in item.get("author", [])]
    editors = [_from_csl_contributor(person) for person in item.get("editor", [])]
    return {
        **({"id": str(item["id"])} if item.get("id") else {}),
        "source_type": supported_types[source_type],
        "title": item.get("title") or "",
        "subtitle": item.get("subtitle"),
        "short_title": item.get("shortTitle"),
        "authors": authors,
        "editors": editors,
        "publication_date": pubdate,
        "journal_title": item.get("container-title") if source_type == "article-journal" else None,
        "book_title": item.get("container-title") if source_type == "chapter" else None,
        "publisher": item.get("publisher"),
        "publisher_place": item.get("publisher-place"),
        "institution": item.get("institution"),
        "volume": item.get("volume"),
        "issue": item.get("issue"),
        "pages": item.get("page"),
        "article_number": item.get("article-number"),
        "edition": item.get("edition"),
        "doi": item.get("DOI"),
        "isbn": item.get("ISBN"),
        "issn": item.get("ISSN"),
        "url": item.get("URL"),
        "language": item.get("language"),
        "number": item.get("number"),
    }


def _from_csl_contributor(person: Any) -> dict[str, str]:
    if not isinstance(person, dict):
        raise FormatError("CSL contributor must be an object.")
    if person.get("literal"):
        return {"literal": str(person["literal"])}
    return {
        "given": str(person.get("given") or ""),
        "family": str(person.get("family") or ""),
    }


def _parse_ris(content: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    type_map = {
        "JOUR": "article-journal",
        "JFULL": "article-journal",
        "BOOK": "book",
        "CHAP": "chapter",
        "THES": "thesis",
        "RPRT": "report",
        "ELEC": "webpage",
        "WEB": "webpage",
    }
    for line in content.splitlines():
        match = re.match(r"^([A-Z0-9]{2})\s*-\s?(.*)$", line)
        if not match:
            continue
        tag, value = match.groups()
        if tag == "TY":
            if value not in type_map:
                raise FormatError(f"Unsupported RIS reference type: {value}")
            current = {
                "source_type": type_map[value],
                "title": "",
                "authors": [],
                "editors": [],
            }
        elif tag == "ER":
            if current:
                records.append(current)
                current = None
        elif current is not None:
            _assign_ris_field(current, tag, value)
    if current:
        records.append(current)
    if not records:
        raise FormatError("No RIS records found.")
    return records


def _assign_ris_field(row: dict[str, Any], tag: str, value: str) -> None:
    simple_fields = {
        "TI": "title",
        "T1": "title",
        "T2": "journal_title",
        "JO": "journal_title",
        "JF": "journal_title",
        "BT": "book_title",
        "PB": "publisher",
        "CY": "publisher_place",
        "VL": "volume",
        "IS": "issue",
        "SP": "pages",
        "DO": "doi",
        "SN": "isbn",
        "UR": "url",
        "AB": "abstract",
        "ET": "edition",
        "M3": "number",
        "PY": "publication_date",
        "Y1": "publication_date",
        "DA": "accessed_date",
        "AD": "institution",
    }
    if tag in {"AU", "A1"}:
        row["authors"].append(_parse_name(value))
    elif tag in {"ED", "A2"}:
        row["editors"].append(_parse_name(value))
    elif tag in simple_fields:
        field = simple_fields[tag]
        row[field] = (
            _normalize_import_date(value)
            if field in {"publication_date", "accessed_date"}
            else value
        )


def _parse_name(value: str) -> dict[str, str]:
    value = value.strip()
    if "," in value:
        family, given = value.split(",", 1)
        return {"family": family.strip(), "given": given.strip()}
    parts = value.split()
    if len(parts) <= 1:
        return {"literal": value}
    return {"given": " ".join(parts[:-1]), "family": parts[-1]}


def _normalize_import_date(value: str) -> str:
    match = re.fullmatch(r"(\d{4})[/.-](\d{1,2})(?:[/.-](\d{1,2}))?", value.strip())
    if match:
        year, month, day = match.groups()
        return f"{year}-{int(month):02d}" + (f"-{int(day):02d}" if day else "")
    return value.strip()


def _parse_bibtex(content: str) -> list[dict[str, Any]]:
    parser = BibTexParser(common_strings=True)
    database = load_bibtex(content, parser=parser)
    type_map = {
        "article": "article-journal",
        "book": "book",
        "inbook": "chapter",
        "incollection": "chapter",
        "inproceedings": "article-journal",
        "misc": "webpage",
        "online": "webpage",
        "techreport": "report",
        "report": "report",
        "phdthesis": "thesis",
        "mastersthesis": "thesis",
        "thesis": "thesis",
    }
    rows: list[dict[str, Any]] = []
    for entry in database.entries:
        entry_type = str(entry.get("ENTRYTYPE", "")).lower()
        if entry_type not in type_map:
            raise FormatError(f"Unsupported BibTeX entry type: {entry_type}")
        container = entry.get("journal") or entry.get("booktitle")
        authors = [
            _parse_bibtex_name(name)
            for name in entry.get("author", "").split(" and ")
            if name.strip()
        ]
        rows.append(
            {
                **({"id": entry["ID"]} if entry.get("ID") else {}),
                "source_type": type_map[entry_type],
                "title": entry.get("title", ""),
                "subtitle": entry.get("subtitle"),
                "authors": authors,
                "editors": [
                    _parse_bibtex_name(name)
                    for name in entry.get("editor", "").split(" and ")
                    if name.strip()
                ],
                "publication_date": entry.get("year"),
                "journal_title": container if entry_type == "article" else None,
                "book_title": container if entry_type in {"inbook", "incollection"} else None,
                "publisher": entry.get("publisher"),
                "publisher_place": entry.get("address"),
                "institution": entry.get("school") or entry.get("institution"),
                "volume": entry.get("volume"),
                "issue": entry.get("number"),
                "pages": entry.get("pages"),
                "edition": entry.get("edition"),
                "doi": entry.get("doi"),
                "isbn": entry.get("isbn"),
                "issn": entry.get("issn"),
                "url": entry.get("url"),
                "number": entry.get("number"),
            }
        )
    if not rows:
        raise FormatError("No BibTeX entries found.")
    return rows


def _parse_bibtex_name(value: str) -> dict[str, str]:
    parts = value.strip().split(",", 1)
    if len(parts) == 2:
        return {"family": parts[0].strip(), "given": parts[1].strip()}
    names = value.split()
    if len(names) <= 1:
        return {"literal": value.strip()}
    return {"given": " ".join(names[:-1]), "family": names[-1]}


def _parse_csv(content: str) -> list[dict[str, Any]]:
    reader = csv.DictReader(io.StringIO(content))
    if (
        not reader.fieldnames
        or "title" not in reader.fieldnames
        or "source_type" not in reader.fieldnames
    ):
        raise FormatError("CSV requires title and source_type columns.")
    rows: list[dict[str, Any]] = []
    for values in reader:
        authors = [
            _parse_name(name) for name in (values.get("authors") or "").split(";") if name.strip()
        ]
        row = {
            key: value or None
            for key, value in values.items()
            if key in Reference.model_fields and key not in {"authors", "editors"}
        }
        year_value = row.get("year")
        if isinstance(year_value, str):
            row["year"] = int(year_value) if year_value else None
        row["authors"] = authors
        row["editors"] = [
            _parse_name(name) for name in (values.get("editors") or "").split(";") if name.strip()
        ]
        rows.append(row)
    if not rows:
        raise FormatError("CSV contains no reference rows.")
    return rows


def _export_csv(references: list[Reference]) -> str:
    buffer = io.StringIO(newline="")
    fields = [
        "id",
        "source_type",
        "title",
        "subtitle",
        "short_title",
        "authors",
        "editors",
        "publication_date",
        "year",
        "journal_title",
        "book_title",
        "publisher",
        "publisher_place",
        "institution",
        "volume",
        "issue",
        "pages",
        "article_number",
        "edition",
        "doi",
        "isbn",
        "issn",
        "url",
        "accessed_date",
        "abstract",
        "language",
        "number",
    ]
    writer = csv.DictWriter(buffer, fieldnames=fields)
    writer.writeheader()
    for reference in references:
        row = reference.model_dump(mode="json")
        for role in ("authors", "editors"):
            row[role] = "; ".join(_display_name(person) for person in getattr(reference, role))
        writer.writerow(row)
    return buffer.getvalue()


def _export_bibtex(references: list[Reference]) -> str:
    type_map = {
        "article-journal": "article",
        "book": "book",
        "chapter": "incollection",
        "webpage": "misc",
        "report": "techreport",
        "thesis": "mastersthesis",
    }
    chunks: list[str] = []
    for index, reference in enumerate(references, start=1):
        key = _bibtex_key(reference, index)
        fields: dict[str, str | None] = {
            "title": reference.title,
            "author": " and ".join(_bibtex_name(person) for person in reference.authors) or None,
            "editor": " and ".join(_bibtex_name(person) for person in reference.editors) or None,
            "year": str(reference.year) if reference.year else None,
            "journal": reference.journal_title,
            "booktitle": reference.book_title,
            "publisher": reference.publisher,
            "address": reference.publisher_place,
            "volume": reference.volume,
            "number": reference.issue or reference.number,
            "pages": reference.pages,
            "edition": reference.edition,
            "doi": reference.doi,
            "isbn": reference.isbn,
            "issn": reference.issn,
            "url": reference.url,
        }
        lines = [f"@{type_map[reference.source_type]}{{{key},"]
        lines.extend(
            f"  {field} = {{{_escape_bibtex(value)}}},"
            for field, value in fields.items()
            if value is not None
        )
        lines.append("}")
        chunks.append("\n".join(lines))
    return "\n\n".join(chunks)


def _export_ris(references: list[Reference]) -> str:
    type_map = {
        "article-journal": "JOUR",
        "book": "BOOK",
        "chapter": "CHAP",
        "webpage": "ELEC",
        "report": "RPRT",
        "thesis": "THES",
    }
    blocks: list[str] = []
    for reference in references:
        lines = [f"TY  - {type_map[reference.source_type]}", f"TI  - {reference.title}"]
        if reference.journal_title:
            lines.append(f"JO  - {reference.journal_title}")
        if reference.book_title:
            lines.append(f"BT  - {reference.book_title}")
        for person in reference.authors:
            lines.append(f"AU  - {_display_name(person)}")
        for person in reference.editors:
            lines.append(f"ED  - {_display_name(person)}")
        mapping = [
            ("publication_date", "PY"),
            ("publisher", "PB"),
            ("publisher_place", "CY"),
            ("volume", "VL"),
            ("issue", "IS"),
            ("pages", "SP"),
            ("doi", "DO"),
            ("isbn", "SN"),
            ("url", "UR"),
            ("abstract", "AB"),
            ("number", "M3"),
            ("institution", "AD"),
            ("edition", "ET"),
        ]
        for field, tag in mapping:
            value = getattr(reference, field)
            if value:
                lines.append(f"{tag}  - {_single_line(value)}")
        lines.append("ER  - ")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def _export_docx(entries: list[str]) -> ExportedFile:
    document = Document()
    document.add_heading("Bibliography", level=1)
    for entry in entries:
        document.add_paragraph(entry, style="List Number")
    buffer = io.BytesIO()
    document.save(buffer)
    return ExportedFile(
        "bibliography.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        base64.b64encode(buffer.getvalue()).decode("ascii"),
        "base64",
    )


def _bibtex_key(reference: Reference, index: int) -> str:
    author = reference.authors[0] if reference.authors else None
    family = (
        author.family if author and author.family else author.literal if author else "reference"
    )
    slug = re.sub(r"[^A-Za-z0-9]+", "", family or "reference").lower() or "reference"
    title_word = re.sub(r"[^A-Za-z0-9]+", "", reference.title.split()[0]).lower()
    return f"{slug}{reference.year or 'nd'}{title_word}{index}"


def _bibtex_name(person: Contributor) -> str:
    if person.literal:
        return "{" + person.literal + "}"
    return f"{person.family or ''}, {person.given or ''}".strip(", ")


def _display_name(person: Contributor) -> str:
    return person.literal or ", ".join(value for value in (person.family, person.given) if value)


def _escape_bibtex(value: str) -> str:
    return value.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}").replace("%", "\\%")


def _single_line(value: str) -> str:
    return re.sub(r"[\r\n]+", " ", value).strip()
